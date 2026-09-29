from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from mentor_feedback_questions import MENTOR_FEEDBACK_OPEN_QUESTIONS


DARK = "263238"
WHITE = "FFFFFF"
PALE = "F3F5F6"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(row):
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _body(row):
    for cell in row:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _fit(sheet, widths):
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35
    sheet.print_title_rows = "1:1"


def export_mentor_feedback_excel(records, feedback_id=None):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Отзывы"
    summary.append([
        "№", "Дата", "Наставник", "Предприятие", "Город", "Участок",
        "Модуль", "Средняя оценка", "Студент (только администратору)",
        "Поток", "Специальность", "Курс",
    ])
    _header(summary[1])
    for record in records:
        summary.append([
            record["id"], record["submitted_at"], record["mentor_name"],
            record["enterprise"], record["city"], record["workshop"],
            record["module"], record["average"], record["student_name"],
            record["stream"], record["speciality"], record["course"],
        ])
        _body(summary[summary.max_row])
    _fit(summary, {"A": 8, "B": 20, "C": 32, "D": 20, "E": 18, "F": 24, "G": 34, "H": 16, "I": 34, "J": 12, "K": 32, "L": 10})

    ratings = workbook.create_sheet("Оценки")
    ratings.append(["№ отзыва", "Дата", "Наставник", "Студент", "Код", "Критерий", "Оценка"])
    _header(ratings[1])
    for record in records:
        for question, score in zip(record["questions"], record["ratings"]):
            ratings.append([
                record["id"], record["submitted_at"], record["mentor_name"],
                record["student_name"], question.get("code", ""),
                question.get("text", ""), int(score),
            ])
            _body(ratings[ratings.max_row])
    _fit(ratings, {"A": 12, "B": 20, "C": 32, "D": 34, "E": 12, "F": 88, "G": 12})

    comments = workbook.create_sheet("Комментарии")
    comments.append(["№ отзыва", "Дата", "Наставник", "Студент", "Раздел", "Ответ"])
    _header(comments[1])
    for record in records:
        for key, title in MENTOR_FEEDBACK_OPEN_QUESTIONS:
            comments.append([
                record["id"], record["submitted_at"], record["mentor_name"],
                record["student_name"], title, record.get(key, ""),
            ])
            _body(comments[comments.max_row])
    _fit(comments, {"A": 12, "B": 20, "C": 32, "D": 34, "E": 58, "F": 82})

    statistics = workbook.create_sheet("Статистика")
    statistics.append(["Наставник", "Предприятие", "Количество отзывов", "Средняя оценка", "Последний отзыв"])
    _header(statistics[1])
    grouped = {}
    for record in records:
        item = grouped.setdefault(
            record["mentor_lms_id"],
            {"name": record["mentor_name"], "enterprise": record["enterprise"], "scores": [], "last": ""},
        )
        item["scores"].append(float(record["average"]))
        item["last"] = max(item["last"], record["submitted_at"])
    for item in sorted(grouped.values(), key=lambda value: value["name"]):
        statistics.append([
            item["name"], item["enterprise"], len(item["scores"]),
            round(sum(item["scores"]) / len(item["scores"]), 2), item["last"],
        ])
        _body(statistics[statistics.max_row])
    _fit(statistics, {"A": 34, "B": 24, "C": 22, "D": 20, "E": 22})

    filename = (
        f"mentor_feedback_{feedback_id}.xlsx"
        if feedback_id is not None
        else f"mentor_feedback_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    workbook.save(filename)
    return filename
